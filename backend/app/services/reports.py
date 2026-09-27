"""Отчёты: предпросмотр и выгрузка.

Главное правило раздела — файл совпадает с тем, что пользователь видел
(ТЗ, раздел 5.6). Поэтому предпросмотр и выгрузка собираются из одной функции,
а не двумя похожими: разойтись они не могут.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Literal, Sequence

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.timeutil import now_utc, to_site
from app.models import DATASET_REAL, Source
from app.services import analytics
from app.services.analytics import Window, build_metric, resolve_window
from app.services.settings import load_settings

LOGGER = logging.getLogger(__name__)

ReportScope = Literal["canteen", "gate", "source"]

FONT_DIR = Path(__file__).resolve().parents[1] / "assets" / "fonts"

# Шаги детализации табличной части. Меньше интервала агрегации шаг быть не может.
STEPS: dict[str, int] = {
    "minute": 60,
    "quarter": 15 * 60,
    "hour": 3600,
    "day": 24 * 3600,
}
STEP_TITLES = {
    "minute": "1 минута",
    "quarter": "15 минут",
    "hour": "1 час",
    "day": "1 сутки",
}
# Больше этого числа строк отчёт перестаёт быть читаемым, и мы предлагаем шаг крупнее.
MAX_ROWS = 400

COLUMNS: list[tuple[str, str, str]] = [
    ("slot", "Временной слот", ""),
    ("peopleInFrame", "В кадре, ср.", "чел"),
    ("peopleInZone", "В зоне, ср.", "чел"),
    ("queueAvg", "Очередь, ср.", "чел"),
    ("queueMax", "Очередь, пик", "чел"),
    ("entered", "Входы", "чел"),
    ("exited", "Выходы", "чел"),
    ("waitMinutes", "Ожидание", "мин"),
    ("latencyMs", "Задержка", "мс"),
]


@dataclass(frozen=True, slots=True)
class ReportRequest:
    scope: ReportScope
    source_id: int | None
    period: str
    date_from: str | None
    date_to: str | None
    step: str | None


def default_step(window: Window) -> str:
    for name, seconds in (("minute", 60), ("quarter", 900), ("hour", 3600), ("day", 86400)):
        if seconds >= window.step_seconds:
            return name
    return "day"


def resolve_step(window: Window, requested: str | None, aggregation_seconds: int) -> tuple[str, str | None]:
    """Шаг таблицы и, если нужно, подсказка о более крупном шаге."""
    step = requested if requested in STEPS else default_step(window)
    if STEPS[step] < aggregation_seconds:
        step = default_step(window)

    length = (window.end - window.start).total_seconds()
    rows = length / STEPS[step]
    if rows > MAX_ROWS:
        bigger = next(
            (name for name, seconds in sorted(STEPS.items(), key=lambda item: item[1])
             if length / seconds <= MAX_ROWS),
            "day",
        )
        return step, (
            f"При шаге «{STEP_TITLES[step]}» получится около {int(rows)} строк. "
            f"Возьмите шаг «{STEP_TITLES[bigger]}» — читать будет проще."
        )
    return step, None


def _sources(session: Session, request: ReportRequest) -> list[Source]:
    if request.scope == "source" and request.source_id:
        source = session.get(Source, request.source_id)
        return [source] if source else []
    return analytics.scope_sources(session, request.scope)


def _title(request: ReportRequest, sources: Sequence[Source]) -> str:
    if request.scope == "source":
        name = sources[0].name if sources else "источник"
        return f"Отчёт по источнику «{name}»"
    return "Отчёт по столовой" if request.scope == "canteen" else "Отчёт по КПП"


def build_report(session: Session, request: ReportRequest, author: str, dataset: str) -> dict[str, Any]:
    """Собрать отчёт: параметры, сводка, табличная часть и приложение."""
    window = resolve_window(request.period, request.date_from, request.date_to)
    settings = load_settings(session)
    sources = _sources(session, request)
    step, hint = resolve_step(window, request.step, settings.aggregation_seconds)

    totals = analytics.collect_totals(session, window, dataset, sources)
    previous = analytics.collect_totals(
        session,
        Window(window.period, window.previous_start, window.previous_end,
               window.previous_start, window.previous_end, window.step_seconds),
        dataset,
        sources,
    )

    table_window = Window(
        period=window.period,
        start=window.start,
        end=window.end,
        previous_start=window.previous_start,
        previous_end=window.previous_end,
        step_seconds=STEPS[step],
    )
    series = analytics.build_series(
        session,
        table_window,
        dataset,
        sources,
        {
            "peopleInFrame": lambda m: analytics.weighted(m.people_in_frame_avg, m),
            "peopleInZone": lambda m: analytics.weighted(m.people_in_zone_avg, m),
            "queueAvg": lambda m: analytics.weighted(m.queue_avg, m),
            "queueMax": lambda m: func.max(m.queue_max),
            "entered": lambda m: func.sum(m.entered),
            "exited": lambda m: func.sum(m.exited),
            "waitMinutes": lambda m: func.sum(m.wait_sum_seconds)
            / func.nullif(func.sum(m.wait_count), 0)
            / 60.0,
            "latencyMs": lambda m: analytics.weighted(m.latency_ms_avg, m),
        },
    )

    # Колонки, неприменимые к объекту отчёта, заполняются прочерками.
    has_queue = any(
        zone["queueAvg"] > 0 for zone in analytics.zone_comparison(session, window, dataset, sources)
    ) or (totals.queue_max or 0) > 0
    rows = [
        {
            "slot": point["at"],
            **{
                key: (None if (key in ("queueAvg", "queueMax", "waitMinutes") and not has_queue)
                      else point.get(key))
                for key, _, _ in COLUMNS
                if key != "slot"
            },
        }
        for point in series
    ]

    events = analytics.count_events(session, window, dataset, sources)
    from app.api.analytics import _event_cards  # локальный импорт: общий вид карточки события

    return {
        "title": _title(request, sources),
        "parameters": {
            "scope": request.scope,
            "sourceId": request.source_id,
            "sourceName": sources[0].name if request.scope == "source" and sources else None,
            "period": window.period,
            "from": window.start.isoformat(),
            "to": window.end.isoformat(),
            "step": step,
            "stepTitle": STEP_TITLES[step],
        },
        "generatedAt": now_utc().isoformat(),
        "author": author,
        "dataset": dataset,
        "demo": dataset != DATASET_REAL,
        "stepHint": hint,
        "sources": [{"id": source.id, "name": source.name} for source in sources],
        "summary": [
            build_metric("entered", "Входы", float(totals.entered),
                         previous=float(previous.entered), unit="чел", whole=True),
            build_metric("exited", "Выходы", float(totals.exited),
                         previous=float(previous.exited), unit="чел", whole=True),
            build_metric("queue_max", "Пик очереди", totals.queue_max,
                         previous=previous.queue_max, unit="чел",
                         missing_reason=None if totals.queue_max is not None else "зоны очереди нет"),
            build_metric("wait_avg", "Среднее ожидание", totals.wait_avg,
                         previous=previous.wait_avg, unit="мин",
                         missing_reason=None if totals.wait_avg is not None else "завершённых ожиданий нет"),
            build_metric("wait_p95", "Ожидание, 95-й процентиль", totals.wait_p95,
                         previous=previous.wait_p95, unit="мин",
                         missing_reason=None if totals.wait_p95 is not None else "завершённых ожиданий нет"),
            build_metric("events", "Событий за период", float(events), unit="шт", whole=True),
        ],
        "columns": [{"key": key, "title": title, "unit": unit} for key, title, unit in COLUMNS],
        "rows": rows,
        "events": _event_cards(session, dataset, sources, window, limit=100),
        "totals": {
            "rows": len(rows),
            "entered": totals.entered,
            "exited": totals.exited,
        },
    }


# --- выгрузка ---------------------------------------------------------------


def _format_slot(value: str) -> str:
    try:
        moment = datetime.fromisoformat(value)
    except ValueError:
        return value
    return to_site(moment).strftime("%d.%m.%Y %H:%M")


def _cell(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:.1f}".replace(".", ",")
    return str(value)


def _period_title(report: dict[str, Any]) -> str:
    start = _format_slot(report["parameters"]["from"])
    end = _format_slot(report["parameters"]["to"])
    return f"{start} — {end}"


def to_csv(report: dict[str, Any]) -> bytes:
    buffer = io.StringIO()
    # Разделитель «;» и десятичная запятая: иначе русский Excel прочитает неверно.
    writer = csv.writer(buffer, delimiter=";")

    writer.writerow([report["title"]])
    if report["demo"]:
        writer.writerow(["Демо-данные"])
    writer.writerow(["Период", _period_title(report)])
    writer.writerow(["Детализация", report["parameters"]["stepTitle"]])
    writer.writerow(["Сформирован", _format_slot(report["generatedAt"]), report["author"]])
    writer.writerow([])

    writer.writerow(["Сводка"])
    for metric in report["summary"]:
        writer.writerow([metric["label"], _cell(metric["value"]), metric["unit"]])
    writer.writerow([])

    writer.writerow([f"{column['title']}{f', {column['unit']}' if column['unit'] else ''}"
                     for column in report["columns"]])
    for row in report["rows"]:
        writer.writerow(
            [_format_slot(row["slot"])] + [_cell(row[column["key"]]) for column in report["columns"][1:]]
        )

    if report["events"]:
        writer.writerow([])
        writer.writerow(["События за период"])
        writer.writerow(["Начало", "Событие", "Приоритет", "Источник", "Пик", "Порог", "Статус"])
        for event in report["events"]:
            writer.writerow([
                _format_slot(event["startedAt"]),
                event["title"],
                event["severity"],
                event["sourceName"],
                _cell(event["peakValue"]),
                _cell(event["threshold"]),
                event["status"],
            ])

    return ("﻿" + buffer.getvalue()).encode("utf-8")


def to_xlsx(report: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Отчёт"

    sheet.append([report["title"]])
    sheet["A1"].font = Font(bold=True, size=14)
    if report["demo"]:
        sheet.append(["Демо-данные"])
    sheet.append(["Период", _period_title(report)])
    sheet.append(["Детализация", report["parameters"]["stepTitle"]])
    sheet.append(["Сформирован", _format_slot(report["generatedAt"]), report["author"]])
    sheet.append([])

    sheet.append(["Сводка"])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    for metric in report["summary"]:
        # Числа кладём числами, а не текстом: иначе Excel не умеет их считать.
        sheet.append([metric["label"], metric["value"], metric["unit"]])
    sheet.append([])

    header = [f"{column['title']}{f', {column['unit']}' if column['unit'] else ''}"
              for column in report["columns"]]
    sheet.append(header)
    for index in range(1, len(header) + 1):
        sheet.cell(row=sheet.max_row, column=index).font = Font(bold=True)

    for row in report["rows"]:
        sheet.append(
            [_format_slot(row["slot"])] + [row[column["key"]] for column in report["columns"][1:]]
        )

    if report["events"]:
        events_sheet = book.create_sheet("События")
        events_sheet.append(["Начало", "Окончание", "Событие", "Приоритет", "Источник", "Пик", "Порог", "Статус"])
        for event in report["events"]:
            events_sheet.append([
                _format_slot(event["startedAt"]),
                _format_slot(event["endedAt"]) if event["endedAt"] else "идёт",
                event["title"],
                event["severity"],
                event["sourceName"],
                event["peakValue"],
                event["threshold"],
                event["status"],
            ])

    widths = [22, 14, 14, 14, 14, 10, 10, 12, 12]
    for index, width in enumerate(widths[: len(header)], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _register_fonts() -> tuple[str, str]:
    """Шрифт с кириллицей: стандартные шрифты PDF её не содержат."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVuSans", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    return "DejaVuSans", "DejaVuSans-Bold"


def to_pdf(report: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    regular, bold = _register_fonts()
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=report["title"],
    )

    title_style = ParagraphStyle("title", fontName=bold, fontSize=16, leading=20)
    text_style = ParagraphStyle("text", fontName=regular, fontSize=9, leading=12)
    note_style = ParagraphStyle("note", fontName=bold, fontSize=10, leading=13,
                                textColor=colors.HexColor("#9a6b12"))

    story: list[Any] = [Paragraph(report["title"], title_style), Spacer(1, 4 * mm)]
    if report["demo"]:
        story.append(Paragraph("Демо-данные", note_style))
        story.append(Spacer(1, 2 * mm))

    story.append(
        Paragraph(
            f"Период: {_period_title(report)} · детализация: {report['parameters']['stepTitle']}<br/>"
            f"Сформирован: {_format_slot(report['generatedAt'])} · {report['author']}",
            text_style,
        )
    )
    story.append(Spacer(1, 5 * mm))

    summary = [["Показатель", "Значение", "Изменение"]]
    for metric in report["summary"]:
        delta = "нет данных" if metric["deltaPercent"] is None else f"{metric['deltaPercent']:+.0f} %"
        summary.append([metric["label"], f"{_cell(metric['value'])} {metric['unit']}".strip(), delta])

    summary_table = Table(summary, hAlign="LEFT", colWidths=[70 * mm, 40 * mm, 35 * mm])
    summary_table.setStyle(
        TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), bold),
            ("FONTNAME", (0, 1), (-1, -1), regular),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6b7480")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#d9dee7")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
        ])
    )
    story.extend([summary_table, Spacer(1, 6 * mm)])

    header = [f"{column['title']}{f', {column['unit']}' if column['unit'] else ''}"
              for column in report["columns"]]
    table_data = [header]
    for row in report["rows"]:
        table_data.append(
            [_format_slot(row["slot"])] + [_cell(row[column["key"]]) for column in report["columns"][1:]]
        )

    table = Table(table_data, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), bold),
            ("FONTNAME", (0, 1), (-1, -1), regular),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6b7480")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#d9dee7")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e7eaf0")),
        ])
    )
    story.append(table)

    if report["events"]:
        story.append(PageBreak())
        story.append(Paragraph("События за период", title_style))
        story.append(Spacer(1, 4 * mm))
        events_data = [["Начало", "Событие", "Приоритет", "Источник", "Пик / порог", "Статус"]]
        for event in report["events"]:
            events_data.append([
                _format_slot(event["startedAt"]),
                event["title"],
                event["severity"],
                event["sourceName"],
                f"{_cell(event['peakValue'])} / {_cell(event['threshold'])}",
                event["status"],
            ])
        events_table = Table(events_data, repeatRows=1, hAlign="LEFT")
        events_table.setStyle(
            TableStyle([
                ("FONTNAME", (0, 0), (-1, 0), bold),
                ("FONTNAME", (0, 1), (-1, -1), regular),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6b7480")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e7eaf0")),
            ])
        )
        story.append(events_table)

    document.build(story)
    return buffer.getvalue()


def filename(report: dict[str, Any], extension: str) -> str:
    stamp = to_site(datetime.fromisoformat(report["generatedAt"])).strftime("%Y-%m-%d-%H%M")
    scope = report["parameters"]["scope"]
    part = {"canteen": "stolovaya", "gate": "kpp", "source": f"source-{report['parameters']['sourceId']}"}
    demo = "-demo" if report["demo"] else ""
    return f"uniflow-{part.get(scope, 'report')}-{stamp}{demo}.{extension}"
