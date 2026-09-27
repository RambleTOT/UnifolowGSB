"""Тесты отчётов: состав, шаг детализации и выгрузка."""

from datetime import timedelta

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.timeutil import now_utc, start_of_day
from app.models import Base, BucketMinute, Source, Wait
from app.services import reports
from app.services.analytics import resolve_window
from app.services.reports import ReportRequest


@pytest.fixture
def session():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with sessionmaker(bind=engine, expire_on_commit=False)() as active:
        session = active
        session.add(Source(id=1, name="Раздача", scope="canteen"))
        session.add(Source(id=2, name="КПП 1", scope="gate"))

        base = start_of_day(now_utc()) + timedelta(hours=9)
        for minute in range(60):
            moment = base + timedelta(minutes=minute)
            session.add(
                BucketMinute(
                    source_id=1,
                    dataset="real",
                    started_at=moment,
                    samples=720,
                    people_in_frame_avg=8.0,
                    people_in_frame_max=11.0,
                    people_in_zone_avg=5.0,
                    queue_avg=6.0,
                    queue_max=9.0,
                    entered=4,
                    exited=3,
                    wait_sum_seconds=360.0,
                    wait_count=2,
                    fps_avg=12.0,
                    latency_ms_avg=32.0,
                    confidence_avg=0.88,
                    health="online",
                )
            )
            session.add(
                Wait(source_id=1, dataset="real", zone_id="z", ended_at=moment, seconds=180.0)
            )
        session.commit()
        yield session


def build(session, **overrides):
    request = ReportRequest(
        scope=overrides.get("scope", "canteen"),
        source_id=overrides.get("source_id"),
        period=overrides.get("period", "day"),
        date_from=None,
        date_to=None,
        step=overrides.get("step"),
    )
    return reports.build_report(session, request, "Администратор", overrides.get("dataset", "real"))


def test_report_has_title_summary_rows_and_columns(session):
    report = build(session)

    assert report["title"] == "Отчёт по столовой"
    assert report["author"] == "Администратор"
    assert report["demo"] is False
    assert [column["key"] for column in report["columns"]][:3] == ["slot", "peopleInFrame", "peopleInZone"]

    summary = {metric["key"]: metric["value"] for metric in report["summary"]}
    assert summary["entered"] == 240  # 4 входа × 60 минут
    assert summary["exited"] == 180
    assert summary["queue_max"] == 9.0
    assert summary["wait_avg"] == 3.0  # 180 секунд на ожидание


def test_step_respects_aggregation_and_warns_about_long_tables(session):
    # Шаг мельче интервала агрегации невозможен: берётся подходящий по периоду.
    day = resolve_window("day")
    step, _ = reports.resolve_step(day, "minute", aggregation_seconds=3600)
    assert step == "quarter"

    # Выбор пользователя сохраняется, но про длину таблицы он предупреждается.
    fine, warning = reports.resolve_step(day, "minute", aggregation_seconds=5)
    assert fine == "minute"
    assert warning is not None and "строк" in warning

    # Без явного выбора шаг подбирается по периоду.
    month = resolve_window("month")
    auto, hint = reports.resolve_step(month, None, aggregation_seconds=5)
    assert auto == "day"
    assert hint is None


def test_source_report_requires_existing_source(session):
    report = build(session, scope="source", source_id=2)
    assert report["title"] == "Отчёт по источнику «КПП 1»"
    assert report["parameters"]["sourceName"] == "КПП 1"


def test_columns_without_queue_are_dashes(session):
    # У КПП зоны очереди нет: соответствующие колонки заполняются прочерками.
    report = build(session, scope="gate")
    filled = [row for row in report["rows"] if row["queueAvg"] is not None]
    assert filled == []


def test_csv_export_matches_preview(session):
    report = build(session)
    content = reports.to_csv(report).decode("utf-8")

    assert content.startswith("﻿")  # русский Excel ждёт метку кодировки
    assert "Отчёт по столовой" in content
    assert "Сводка" in content
    # Десятичная запятая и разделитель «;» — иначе Excel превратит числа в даты.
    assert ";" in content
    assert "3,0" in content
    assert len([line for line in content.splitlines() if line]) > len(report["rows"])


def test_xlsx_export_keeps_numbers_as_numbers(session):
    import io

    from openpyxl import load_workbook

    report = build(session)
    book = load_workbook(io.BytesIO(reports.to_xlsx(report)))
    sheet = book["Отчёт"]

    values = [cell.value for row in sheet.iter_rows(min_row=1, max_row=12) for cell in row]
    assert "Отчёт по столовой" in values
    assert any(isinstance(value, (int, float)) for value in values)


def test_pdf_export_contains_cyrillic_font(session):
    report = build(session)
    content = reports.to_pdf(report)

    assert content.startswith(b"%PDF")
    # Кириллица в PDF работает только со встроенным шрифтом — проверяем, что он внедрён.
    assert b"DejaVuSans" in content


def test_demo_report_is_marked(session):
    report = build(session, dataset="demo")
    assert report["demo"] is True
    assert "Демо-данные" in reports.to_csv(report).decode("utf-8")
    assert "demo" in reports.filename(report, "pdf")
